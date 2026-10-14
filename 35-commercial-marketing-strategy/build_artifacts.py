#!/usr/bin/env python3
"""Собирает рабочий документ, презентацию, граф и интеллект-карту."""
from __future__ import annotations

import html
import json
import math
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parent
GRAPH = json.loads((ROOT / "strategy_graph.json").read_text(encoding="utf-8"))

COLORS = {
    "Цель": "#f59e0b", "Исходные данные": "#94a3b8", "Производство": "#ef4444",
    "Управление": "#8b5cf6", "Команда": "#3b82f6", "Продажи": "#06b6d4",
    "Маркетинг": "#22c55e", "ИИ": "#ec4899", "2027": "#f97316",
}


def parse_inline(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.strip()


def build_docx() -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    src = (ROOT / "STRATEGY.md").read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.7); sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.5)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"; styles["Normal"].font.size = Pt(10.5)
    for name, size, color in (("Title", 24, "0F172A"), ("Heading 1", 17, "0F3D5E"), ("Heading 2", 14, "155E75"), ("Heading 3", 12, "0F766E")):
        styles[name].font.name = "Arial"; styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)
    styles["Title"].font.bold = True

    i = 0
    while i < len(src):
        line = src[i].rstrip()
        if not line or line == "---":
            i += 1; continue
        if line.startswith("```"):
            code = [] ; i += 1
            while i < len(src) and not src[i].startswith("```"):
                code.append(src[i]); i += 1
            p = doc.add_paragraph("\n".join(code)); p.style = styles["Normal"]
            p.paragraph_format.left_indent = Cm(0.8)
            for run in p.runs: run.font.name = "Courier New"; run.font.size = Pt(9)
            i += 1; continue
        if line.startswith("|") and i + 1 < len(src):
            rows = []
            while i < len(src) and src[i].startswith("|"):
                cells = [parse_inline(x.strip()) for x in src[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", x or " ") for x in cells): rows.append(cells)
                i += 1
            if rows:
                cols = max(map(len, rows)); table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for r, values in enumerate(rows):
                    for c, value in enumerate(values):
                        cell = table.cell(r, c); cell.text = value
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.font.name = "Arial"; run.font.size = Pt(8.5)
                                if r == 0: run.bold = True; run.font.color.rgb = RGBColor.from_string("FFFFFF")
                        if r == 0:
                            cell._tc.get_or_add_tcPr().append(__import__("docx").oxml.parse_xml('<w:shd {} w:fill="0F3D5E"/>'.format(__import__("docx").oxml.ns.nsdecls('w'))))
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(parse_inline(line[2:]), style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(parse_inline(line[3:]), level=1)
        elif line.startswith("### "):
            doc.add_heading(parse_inline(line[4:]), level=2)
        elif line.startswith("#### "):
            doc.add_heading(parse_inline(line[5:]), level=3)
        elif re.match(r"^\d+\. ", line):
            doc.add_paragraph(parse_inline(re.sub(r"^\d+\. ", "", line)), style="List Number")
        elif line.startswith("- "):
            doc.add_paragraph(parse_inline(line[2:]), style="List Bullet")
        elif line.startswith("> "):
            p = doc.add_paragraph(parse_inline(line[2:])); p.paragraph_format.left_indent = Cm(0.8)
            for run in p.runs: run.italic = True; run.font.color.rgb = RGBColor.from_string("0F766E")
        else:
            p = doc.add_paragraph(parse_inline(line))
            p.paragraph_format.space_after = Pt(5)
        i += 1

    footer = sec.footer.paragraphs[0]
    footer.text = "ТОО «Стар Билдинг» / Grand Construction — рабочая версия стратегии 2026–2027"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(ROOT / "star_building_strategy.docx")


def graph_html() -> str:
    data = json.dumps(GRAPH, ensure_ascii=False)
    colors = json.dumps(COLORS, ensure_ascii=False)
    return f'''<!doctype html>
<html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Граф стратегии Star Building</title><style>
:root{{--bg:#06111f;--panel:#0d1b2a;--ink:#e5eef8;--muted:#8aa2b8;--line:#27445d}}
*{{box-sizing:border-box}}body{{margin:0;background:radial-gradient(circle at 20% 0,#12304a,#06111f 55%);color:var(--ink);font:14px Arial,sans-serif;overflow:hidden}}
header{{height:82px;padding:16px 22px;border-bottom:1px solid var(--line);display:flex;align-items:center;gap:22px;background:#071522e8}}
h1{{font-size:20px;margin:0;white-space:nowrap}}.sub{{color:var(--muted);font-size:12px}}input,select,button{{background:#102538;color:var(--ink);border:1px solid #31536e;border-radius:8px;padding:9px 12px}}
button{{cursor:pointer}}#stage{{height:calc(100vh - 82px);display:grid;grid-template-columns:1fr 310px}}#canvas{{position:relative;overflow:hidden}}svg{{width:100%;height:100%;cursor:grab}}svg:active{{cursor:grabbing}}
#detail{{border-left:1px solid var(--line);padding:18px;background:#091827;overflow:auto}}#detail h2{{font-size:18px;margin-top:0}}#detail .tag{{display:inline-block;padding:4px 8px;border-radius:20px;background:#17344b;color:#b9d7ee;margin:3px 4px 3px 0;font-size:11px}}
.legend{{display:flex;flex-wrap:wrap;gap:7px;margin-top:15px}}.legend span{{font-size:11px;padding:4px 7px;border-radius:12px;color:#07111b;font-weight:bold}}
.edge{{stroke:#46677f;stroke-width:1.1;opacity:.52}}.node circle{{stroke:#dbeafe;stroke-width:1;cursor:pointer;filter:drop-shadow(0 4px 6px #0008)}}.node text{{fill:#f8fafc;font-size:10px;pointer-events:none;text-anchor:middle}}.node.selected circle{{stroke:#fff;stroke-width:4}}.edge.active{{stroke:#fbbf24;stroke-width:3;opacity:1}}
.toolbar{{display:flex;gap:8px;align-items:center;flex:1}}#search{{width:240px}}.stats{{margin-left:auto;color:#9ec2dc;font-size:12px}}
@media(max-width:900px){{#stage{{grid-template-columns:1fr}}#detail{{display:none}}header{{height:120px;flex-wrap:wrap}}#stage{{height:calc(100vh - 120px)}}}}
</style></head><body>
<header><div><h1>Граф стратегии Star Building</h1><div class="sub">Продажи · производство · маркетинг · ИИ · Астана · B2G</div></div>
<div class="toolbar"><input id="search" placeholder="Поиск узла"><select id="group"><option value="">Все контуры</option></select><button id="reset">Сбросить</button></div><div class="stats" id="stats"></div></header>
<div id="stage"><div id="canvas"><svg id="svg"><g id="world"><g id="edges"></g><g id="nodes"></g></g></svg></div>
<aside id="detail"><h2>Управляемые 500 т/мес</h2><p>Выберите узел, чтобы увидеть его роль и связи.</p><div class="legend" id="legend"></div><hr style="border-color:#27445d"><p class="sub">Граф автономный: все данные встроены в файл. Колесо — масштаб, перетаскивание — перемещение.</p></aside></div>
<script>
const DATA={data}; const COLORS={colors}; const svg=document.getElementById('svg'),world=document.getElementById('world');
const W=1500,H=900,cx=W/2,cy=H/2; svg.setAttribute('viewBox',`0 0 ${{W}} ${{H}}`);
const groups=[...new Set(DATA.nodes.map(n=>n.group))]; document.getElementById('stats').textContent=`${{DATA.nodes.length}} узлов · ${{DATA.edges.length}} связей`;
const sel=document.getElementById('group'); groups.forEach(g=>{{let o=document.createElement('option');o.value=g;o.textContent=g;sel.appendChild(o)}});
document.getElementById('legend').innerHTML=groups.map(g=>`<span style="background:${{COLORS[g]}}">${{g}}</span>`).join('');
const index=Object.fromEntries(DATA.nodes.map((n,i)=>[n.id,n]));
const buckets=Object.fromEntries(groups.map(g=>[g,DATA.nodes.filter(n=>n.group===g)]));
groups.forEach((g,gi)=>{{const a=-Math.PI/2+gi*2*Math.PI/groups.length, bx=cx+Math.cos(a)*315,by=cy+Math.sin(a)*315; buckets[g].forEach((n,j)=>{{const d=(j-(buckets[g].length-1)/2)*42; n.x=bx+Math.cos(a+Math.PI/2)*d;n.y=by+Math.sin(a+Math.PI/2)*d}})}}); index.goal.x=cx;index.goal.y=cy;
const NS='http://www.w3.org/2000/svg',eg=document.getElementById('edges'),ng=document.getElementById('nodes');
DATA.edges.forEach((e,i)=>{{let a=index[e.source],b=index[e.target],l=document.createElementNS(NS,'line');l.setAttribute('x1',a.x);l.setAttribute('y1',a.y);l.setAttribute('x2',b.x);l.setAttribute('y2',b.y);l.setAttribute('class','edge');l.dataset.i=i;eg.appendChild(l)}});
function short(s){{return s.length>24?s.slice(0,22)+'…':s}}
DATA.nodes.forEach(n=>{{let g=document.createElementNS(NS,'g');g.setAttribute('class','node');g.setAttribute('transform',`translate(${{n.x}},${{n.y}})`);g.dataset.id=n.id;let c=document.createElementNS(NS,'circle');c.setAttribute('r',n.id==='goal'?42:27);c.setAttribute('fill',COLORS[n.group]||'#94a3b8');let t=document.createElementNS(NS,'text');t.setAttribute('y',n.id==='goal'?3:4);t.textContent=short(n.label);g.append(c,t);g.onclick=()=>selectNode(n.id);ng.appendChild(g)}});
function selectNode(id){{document.querySelectorAll('.node').forEach(x=>x.classList.toggle('selected',x.dataset.id===id));document.querySelectorAll('.edge').forEach((x,i)=>{{let e=DATA.edges[i];x.classList.toggle('active',e.source===id||e.target===id)}});let n=index[id], rel=DATA.edges.filter(e=>e.source===id||e.target===id);document.getElementById('detail').innerHTML=`<h2>${{n.label}}</h2><span class="tag">${{n.group}}</span><span class="tag">${{n.type}}</span><p>${{n.description||'Элемент единой стратегии.'}}</p><h3>Связи</h3>${{rel.map(e=>`<p><b>${{e.source===id?'→':'←'}} ${{e.relation}}</b><br>${{index[e.source===id?e.target:e.source].label}}</p>`).join('')}}<hr style="border-color:#27445d"><div class="legend" id="legend2">${{groups.map(g=>`<span style="background:${{COLORS[g]}}">${{g}}</span>`).join('')}}</div>`}}
function filter(){{let q=document.getElementById('search').value.toLowerCase(),gr=sel.value;document.querySelectorAll('.node').forEach(x=>{{let n=index[x.dataset.id],ok=(!gr||n.group===gr)&&(!q||n.label.toLowerCase().includes(q));x.style.opacity=ok?1:.08}})}}
document.getElementById('search').oninput=filter;sel.onchange=filter;document.getElementById('reset').onclick=()=>{{document.getElementById('search').value='';sel.value='';filter();selectNode('goal')}};
let scale=1,tx=0,ty=0,drag=false,px=0,py=0;function tr(){{world.setAttribute('transform',`translate(${{tx}} ${{ty}}) scale(${{scale}})`)}}svg.onwheel=e=>{{e.preventDefault();scale=Math.max(.45,Math.min(2.4,scale*(e.deltaY<0?1.1:.9)));tr()}};svg.onpointerdown=e=>{{drag=true;px=e.clientX;py=e.clientY;svg.setPointerCapture(e.pointerId)}};svg.onpointermove=e=>{{if(!drag)return;tx+=(e.clientX-px)/scale;ty+=(e.clientY-py)/scale;px=e.clientX;py=e.clientY;tr()}};svg.onpointerup=()=>drag=false;selectNode('goal');
</script></body></html>'''


def mindmap_html() -> str:
    branches = {
        "Производство": ["600 т по слотам", "Аудит 10 дней", "Персонал и сменность", "150–200 → 250–300 → 350–400 → 450–500 т", "Качество и оборотный капитал"],
        "Коммерческий отдел": ["Коммерческий директор", "РОП", "Ведущий менеджер", "Координатор Астаны", "Менеджер по тендерам", "Новый B2B-менеджер"],
        "Продажи и Битрикс24": ["Воронка 1000–2000 т", "9–10 побед в месяц", "4 продуктовых пакета", "Слоты и конечные сроки", "Панель тонн, маржи и оплат"],
        "Маркетинг": ["STAR BUILDING — завод", "Grand Construction — под ключ", "Кейсы и экскурсии", "Проектировщики и партнёры", "Просвещение B2B/B2G"],
        "ИИ": ["ChatGPT / Claude", "Midjourney / Recraft / Firefly", "Whisper / CapCut / Runway", "n8n / Make + Битрикс24", "Человеческая проверка"],
        "2027": ["Проверка рынка Астаны", "Лёгкое представительство", "Офис после контрольных рубежей", "Нормативно-технический пакет", "Пилотный B2G-объект"],
    }
    cards = []
    for i, (branch, items) in enumerate(branches.items()):
        color = list(COLORS.values())[(i+2) % len(COLORS)]
        lis = "".join(f"<li><span>{html.escape(x)}</span></li>" for x in items)
        cards.append(f'<section style="--c:{color}"><button class="branch">{html.escape(branch)}<b>−</b></button><ul>{lis}</ul></section>')
    decisions = "".join(f"<article><b>{i+1}</b><span>{html.escape(text)}</span></article>" for i, text in enumerate([
        "Распределить портфель 600 тонн по производственным слотам",
        "Провести десятидневный аудит мощности и узких мест",
        "Утвердить границы Коммерческого директора и РОПа",
        "Нанять производственный персонал и проектного B2B-менеджера",
        "Переходить к следующей ступени мощности только после двух периодов выполнения плана на 90%",
        "Открывать постоянный офис в Астане только после подтверждения спроса и экономики",
    ]))
    return f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Интеллект-карта стратегии</title><style>
*{{box-sizing:border-box}}body{{margin:0;background:#07131f;color:#eef7ff;font:16px Arial,sans-serif;min-height:100vh}}header{{padding:26px 30px;text-align:center;background:linear-gradient(135deg,#0f3d5e,#0f766e)}}h1{{margin:0;font-size:28px}}header p{{margin:8px;color:#d6f4ff}}main{{position:relative;max-width:1450px;margin:auto;padding:38px 30px 70px}}.core{{width:330px;margin:0 auto 40px;padding:28px;border:3px solid #fbbf24;border-radius:50%;background:#172554;text-align:center;box-shadow:0 12px 40px #0008}}.core strong{{font-size:26px;color:#fde68a}}.core small{{display:block;margin-top:7px;color:#bdd7ea}}.grid{{display:grid;grid-template-columns:repeat(3,1fr);gap:28px;position:relative}}section{{border:2px solid var(--c);border-radius:16px;background:#0d2031;overflow:hidden;box-shadow:0 10px 25px #0005;position:relative}}section:before{{content:"";position:absolute;top:-30px;left:50%;height:30px;border-left:2px solid var(--c)}}.branch{{width:100%;border:0;background:var(--c);color:#07111b;padding:15px 16px;text-align:left;font-size:18px;font-weight:bold;cursor:pointer}}.branch b{{float:right}}ul{{margin:0;padding:16px 22px 18px 42px;transition:.2s}}li{{padding:7px 0;color:#d5e8f7}}li::marker{{color:var(--c)}}section.closed ul{{display:none}}.flow{{margin:42px auto 0;max-width:1100px;display:flex;align-items:center;justify-content:center;flex-wrap:wrap;gap:8px}}.flow span{{padding:10px 12px;border:1px solid #3d627e;border-radius:9px;background:#10283c}}.flow i{{color:#fbbf24;font-style:normal;font-size:20px}}.decisions{{margin:45px auto 0;max-width:1100px}}.decisions h2{{text-align:center;color:#fde68a}}.decision-grid{{display:grid;grid-template-columns:1fr 1fr;gap:12px}}article{{display:flex;gap:14px;align-items:center;padding:14px;border:1px solid #34536b;border-radius:12px;background:#0b1d2b}}article b{{display:grid;place-items:center;min-width:34px;height:34px;border-radius:50%;background:#f59e0b;color:#07111b}}article span{{line-height:1.35;color:#d9edf9}}footer{{text-align:center;color:#7896ad;padding:20px}}@media(max-width:950px){{.grid{{grid-template-columns:1fr 1fr}}}}@media(max-width:640px){{.grid,.decision-grid{{grid-template-columns:1fr}}.core{{width:280px}}}}
</style></head><body><header><h1>Интеллект-карта стратегии Star Building</h1><p>Единая система коммерции, производства, маркетинга и масштабирования</p></header><main><div class="core"><strong>500 тонн в месяц</strong><small>Продать · произвести · отгрузить · получить оплату</small></div><div class="grid">{''.join(cards)}</div><div class="flow"><span>Спрос</span><i>→</i><span>Договор</span><i>→</i><span>Подтверждённый слот</span><i>→</i><span>Производство</span><i>→</i><span>Отгрузка</span><i>→</i><span>Оплата</span></div><div class="decisions"><h2>Ключевые решения руководства</h2><div class="decision-grid">{decisions}</div><h2>Контрольные показатели</h2><div class="decision-grid"><article><b>1</b><span><strong>Портфель:</strong> 100% договоров на 600 тонн имеют приоритет, подтверждённый слот, ответственного и риск.</span></article><article><b>2</b><span><strong>Воронка:</strong> 1000–2000 квалифицированных тонн с вероятностью, сроком и следующим действием.</span></article><article><b>3</b><span><strong>Продажи:</strong> около 9–10 побед в месяц при среднем заказе 50–60 тонн после выхода производства на целевую мощность.</span></article><article><b>4</b><span><strong>Производство:</strong> трудочасы на тонну, простои, брак, своевременность выпуска и отгрузки.</span></article><article><b>5</b><span><strong>Маркетинг:</strong> целевые компании и квалифицированные тонны вместо общего охвата и числа подписчиков.</span></article><article><b>6</b><span><strong>Астана:</strong> 200 организаций, 30 приоритетных заказчиков, 5 партнёров и воронка не менее 1500 тонн до постоянного офиса.</span></article></div><h2>Управленческий ритм</h2><div class="decision-grid"><article><b>13</b><span>Тринадцать недель — общий прогноз загрузки, продаж, выпуска, отгрузки и денежных поступлений.</span></article><article><b>8</b><span>Восемь недель — предварительное резервирование производственных мощностей и материалов.</span></article><article><b>4</b><span>Четыре недели — утверждённый производственный план с объектами и ответственными.</span></article><article><b>1</b><span>Одна неделя — сменные задания и ежедневный план/факт по участкам.</span></article></div></div></main><footer>Рабочая версия 2026–2027 · нажмите на заголовок ветви, чтобы свернуть её</footer><script>document.querySelectorAll('.branch').forEach(b=>b.onclick=()=>{{let s=b.parentElement;s.classList.toggle('closed');b.querySelector('b').textContent=s.classList.contains('closed')?'+':'−'}})</script></body></html>'''


def build_excalidraw() -> None:
    branches = [
        ("Производство", "600 т по слотам\nАудит мощности\nПерсонал и сменность", "#ffc9c9"),
        ("Коммерческий отдел", "Комдир / РОП\nВедущий и новый менеджер\nТендеры / Астана", "#a5d8ff"),
        ("Продажи", "Битрикс24\nВоронка 1000–2000 т\n9–10 побед/мес", "#c3fae8"),
        ("Маркетинг", "Два бренда\nКейсы и партнёры\nB2B/B2G-просвещение", "#b2f2bb"),
        ("ИИ", "ChatGPT / Claude\nMidjourney / Recraft\nWhisper / видео / n8n", "#d0bfff"),
        ("2027", "Астана\nГосбазы и тендеры\nПилотный объект", "#ffd8a8"),
    ]
    elements = []
    def shape_text(id_, x,y,w,h,text,color,ellipse=False):
        typ = "ellipse" if ellipse else "rectangle"
        elements.append({"type":typ,"id":id_,"x":x,"y":y,"width":w,"height":h,"strokeColor":"#1e1e1e","backgroundColor":color,"fillStyle":"solid","strokeWidth":2,"roughness":1,"opacity":100,"roundness":{"type":3} if not ellipse else None,"boundElements":[{"id":"t_"+id_,"type":"text"}]})
        elements.append({"type":"text","id":"t_"+id_,"x":x+10,"y":y+10,"width":w-20,"height":h-20,"text":text,"fontSize":20 if ellipse else 17,"fontFamily":1,"strokeColor":"#1e1e1e","textAlign":"center","verticalAlign":"middle","containerId":id_,"originalText":text,"autoResize":True})
    shape_text("core",700,430,330,150,"500 тонн в месяц\nЕдиная коммерческо-производственная система","#fff3bf",True)
    positions=[(80,80),(700,40),(1320,80),(80,740),(700,800),(1320,740)]
    for i,((title,body,color),(x,y)) in enumerate(zip(branches,positions)):
        id_=f"b{i}"; shape_text(id_,x,y,360,175,title+"\n\n"+body,color)
        sx=865;sy=505;ex=x+180;ey=y+87
        elements.append({"type":"arrow","id":f"a{i}","x":sx,"y":sy,"width":ex-sx,"height":ey-sy,"points":[[0,0],[ex-sx,ey-sy]],"strokeColor":"#495057","strokeWidth":2,"roughness":1,"opacity":100,"endArrowhead":"arrow","startBinding":{"elementId":"core","fixedPoint":[0.5,0.5]},"endBinding":{"elementId":id_,"fixedPoint":[0.5,0.5]}})
    flow=["Спрос","Договор","Слот","Производство","Отгрузка","Оплата"]
    for i,label in enumerate(flow):
        x=255+i*240;y=1080;shape_text(f"f{i}",x,y,160,70,label,"#e7f5ff")
        if i:
            elements.append({"type":"arrow","id":f"fa{i}","x":x-80,"y":1115,"width":80,"height":0,"points":[[0,0],[80,0]],"strokeColor":"#0f766e","strokeWidth":2,"roughness":1,"opacity":100,"endArrowhead":"arrow"})
    controls=["Аудит\n10 дней","Комитет\nеженедельно","План\n13/8/4/1","Качество\n≥90% плана","Воронка\nдо 2000 т","Астана\n1500 т","B2G\nтехпакет","Панель\nтонн и маржи"]
    for i,label in enumerate(controls):
        x=110+i*205;y=1270;shape_text(f"c{i}",x,y,170,82,label,"#e7f5ff")
    out={"type":"excalidraw","version":2,"source":"hermes-agent","elements":elements,"appState":{"viewBackgroundColor":"#ffffff"}}
    (ROOT / "strategy_mindmap.excalidraw").write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding="utf-8")


def add_textbox(slide, text, left, top, width, height, size=18, color="E2E8F0", bold=False):
    from pptx.util import Inches, Pt
    box=slide.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height));tf=box.text_frame;tf.clear();tf.word_wrap=True
    p=tf.paragraphs[0];p.text=text;p.font.size=Pt(size);p.font.name="Arial";p.font.bold=bold;p.font.color.rgb=__import__('pptx').dml.color.RGBColor.from_string(color)
    return box


def build_pptx() -> None:
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    bg="07131F"; ink="EAF4FB"; muted="9EC1D9"; accent="F59E0B"; teal="14B8A6"
    def slide(title,subtitle=None):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(bg)
        add_textbox(s,title,.55,.35,12.2,.55,24,ink,True)
        if subtitle:add_textbox(s,subtitle,.58,.92,12,.35,11,muted)
        shape=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(.55),Inches(1.3),Inches(1.4),Inches(.06));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(accent);shape.line.fill.background()
        add_textbox(s,"STAR BUILDING · GRAND CONSTRUCTION",.58,7.08,6,.2,8,"6F91A8")
        return s
    def bullets(s,items,left=.7,top=1.65,width=12,height=4.9,size=19):
        box=s.shapes.add_textbox(Inches(left),Inches(top),Inches(width),Inches(height));tf=box.text_frame;tf.clear();tf.word_wrap=True
        for i,item in enumerate(items):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=item;p.font.name="Arial";p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(ink);p.space_after=Pt(12);p.level=0
        return box
    s=slide("Единая коммерческо-производственная стратегия","Рабочая версия для решения Генерального директора · 2026–2027");add_textbox(s,"500 тонн в месяц",.75,2.0,11.8,1.1,44,"FBBF24",True);add_textbox(s,"Продать · произвести · отгрузить · получить оплату",.8,3.15,11.5,.6,24,"CDE8F6")
    s=slide("Исходная точка: цель выше текущего темпа примерно в 6 раз");
    chart_data=ChartData();chart_data.categories=["Средний месяц 2025","Целевой месяц"];chart_data.add_series("тонн",(83,500));chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(.7),Inches(1.65),Inches(6),Inches(4.6),chart_data).chart;chart.has_legend=False;chart.value_axis.maximum_scale=550;chart.value_axis.tick_labels.font.color.rgb=RGBColor.from_string(ink);chart.category_axis.tick_labels.font.color.rgb=RGBColor.from_string(ink);chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=RGBColor.from_string(teal)
    bullets(s,["Около 1000 тонн произведено за 2025 год","Около 600 тонн договоров без единого графика","Средний заказ 50–60 тонн","Цикл сделки 2–3 месяца"],7.25,1.8,5.3,4.7,18)
    s=slide("Главное ограничение — не спрос, а исполнимая мощность");bullets(s,["Оборудование имеется, но необходимы наладка, персонал и сменность","Каждый договор должен получить подтверждённый производственный слот","Обещание срока клиенту — только после подтверждения производства","Цель: рост без потери качества, сроков и оборотного капитала"],.85,1.7,11.7,4.8,21)
    s=slide("Единая цепочка управления");steps=["Спрос","Квалификация","Договор","Слот","Производство","Отгрузка","Оплата"]
    for i,x in enumerate(steps):
        left=.45+i*1.82;sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(left),Inches(2.7),Inches(1.45),Inches(1));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string("0F766E" if i not in (3,4) else "B45309");sh.line.color.rgb=RGBColor.from_string("DDF6FF");tf=sh.text_frame;tf.text=x;tf.paragraphs[0].font.size=Pt(15);tf.paragraphs[0].font.bold=True;tf.paragraphs[0].font.color.rgb=RGBColor.from_string("FFFFFF");tf.paragraphs[0].alignment=1
        if i<6:add_textbox(s,"→",left+1.48,2.97,.3,.3,18,"FBBF24",True)
    s=slide("Команда: роли без дублирования");bullets(s,["Коммерческий директор — рынок, маржа, каналы и согласование с мощностью","РОП — ежедневная дисциплина, обучение, воронка и прогноз","Ведущий менеджер — крупные сделки A/B","Координатор Астаны — рынок, встречи и объекты под наставничеством","Менеджер по тендерам — отдельная воронка и фильтр участия","Новый B2B-менеджер — системное пополнение проектной воронки"],.75,1.55,12,5.15,18)
    s=slide("Производственный рост — только ступенями");vals=[("150–200","Стабилизация"),("250–300","Первая смена"),("350–400","Сменность и материалы"),("450–500","Целевая мощность")]
    for i,(v,lbl) in enumerate(vals):
        left=.65+i*3.15;add_textbox(s,v+" т/мес",left,2.0,2.7,.65,25,"FBBF24",True);add_textbox(s,lbl,left,2.72,2.7,.6,15,ink);add_textbox(s,"Переход после двух периодов ≥90% плана" if i<3 else "Люди + капитал + загрузка",left,3.55,2.65,1.05,13,muted)
    s=slide("Битрикс24: управление тоннами, сроками и маржой");bullets(s,["Воронки: спрос, проектная продажа, исполнение, тендеры и партнёры","Обязательные поля: объект, тоннаж, ЛПР, бюджет, срок, следующий шаг","Сделка связывается с проектированием, материалом и производственным слотом","Панель: квалифицировано → законтрактовано → принято → выпущено → отгружено → оплачено"],.75,1.65,12,4.9,20)
    s=slide("Коммерческая модель 500 тонн");add_textbox(s,"≈ 9",.8,1.9,2.2,1,45,"FBBF24",True);add_textbox(s,"побед в месяц\nпри заказе 55 т",.8,3.0,2.4,1.2,17,ink);add_textbox(s,"1000–2000 т",4.0,1.9,3.5,1,38,"22D3EE",True);add_textbox(s,"квалифицированная воронка\nпри цикле 2–3 месяца",4.0,3.0,3.4,1.2,17,ink);add_textbox(s,"4 пакета",8.7,1.9,3,1,38,"4ADE80",True);add_textbox(s,"каркас · тёплый контур\nбыстрый объект · серия",8.7,3.0,3.4,1.2,17,ink)
    s=slide("Маркетинг: не охваты, а квалифицированные тонны");bullets(s,["STAR BUILDING — завод, технология и производство","Grand Construction — объект под ключ","Кейсы, экскурсии, инженерные разборы и партнёрства","Проектировщики и подрядчики — источник раннего входа в проекты","B2B/B2G-просвещение формирует категорию ЛСТК"],.8,1.55,11.8,5.1,21)
    s=slide("ИИ-контур вместо ручного производства графики");bullets(s,["ChatGPT / Claude — тексты, сценарии и протоколы","Midjourney / Recraft / Firefly — визуальные концепции и графика","Whisper / CapCut / Runway / Kling / Veo — интервью и видео","n8n / Make — передача материалов и результатов в Битрикс24","Человек проверяет факты, технику и маркировку концептуальных материалов"],.8,1.55,11.8,5.1,20)
    s=slide("2027: Астана и системный выход в B2G");bullets(s,["Астана: 200 организаций → 30 приоритетов → 5 партнёров → воронка 1500 т","Сначала лёгкое представительство; постоянный офис — после контрольных рубежей","B2G: карта барьеров, испытания, расчёты, узлы, сметные ресурсы","Пилотный объект превращает технологию в измеримое доказательство"],.8,1.6,11.8,4.9,20)
    s=slide("Решения Генерального директора");bullets(s,["Утвердить коммерческо-производственный комитет","Утвердить аудит мощности и слоты для 600 тонн","Разрешить подбор производственного персонала и B2B-менеджера","Зафиксировать границы Коммерческого директора и РОПа","Утвердить ступенчатый выход на 500 тонн","Утвердить проверку рынка Астаны и B2G-проект на 2027 год"],.8,1.5,11.8,5.2,19)
    prs.save(ROOT / "star_building_strategy_presentation.pptx")


def main() -> None:
    build_docx()
    (ROOT / "strategy_graph.html").write_text(graph_html(), encoding="utf-8")
    (ROOT / "strategy_mindmap.html").write_text(mindmap_html(), encoding="utf-8")
    build_excalidraw()
    build_pptx()
    print(json.dumps({"status":"ok","outputs":[p.name for p in ROOT.iterdir() if p.suffix in {'.docx','.pptx','.html','.excalidraw'}]},ensure_ascii=False))


if __name__ == "__main__":
    main()
