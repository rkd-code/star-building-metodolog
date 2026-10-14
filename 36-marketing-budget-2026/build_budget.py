#!/usr/bin/env python3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from budget_model import ALLOCATION, CATEGORIES, KPI_TARGETS, MANDATORY_INFRA, MONTHLY_OPERATING, MONTHS, ONE_TIME_CAPEX, monthly_total, totals

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "marketing_budget_sep_dec_2026.xlsx"
DOCX = ROOT / "marketing_budget_approval.docx"

NAVY = "123047"
TEAL = "0F766E"
BLUE = "DDEBF7"
LIGHT = "F3F6F8"
GOLD = "F4B183"
GREEN = "E2F0D9"
RED = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "64748B"
MONEY = '#,##0 "₸"'
PERCENT = "0.0%"
THIN = Side(style="thin", color="CBD5E1")


def title(ws, text, end_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, text)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.font = Font(color=WHITE, bold=True, size=16)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28


def header(row):
    for c in row:
        c.fill = PatternFill("solid", fgColor=TEAL)
        c.font = Font(color=WHITE, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=THIN)


def body_style(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for c in row:
            c.border = Border(bottom=Side(style="hair", color="D7DEE5"))
            c.alignment = Alignment(vertical="top", wrap_text=True)


def add_table(ws, ref, name):
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)


def build_dashboard(wb):
    ws = wb.active
    ws.title = "Панель"
    title(ws, "Маркетинговый бюджет Star Building · сентябрь–декабрь 2026", 8)
    ws["A3"] = "Показатель"; ws["B3"] = "План"
    header(ws[3])
    values = [
        ("Операционный маркетинг", "=SUM('Помесячный бюджет'!B5:B8)"),
        ("Обязательная инфраструктура", "=SUM('Помесячный бюджет'!C5:C8)"),
        ("Разовое оборудование", "=SUM('Помесячный бюджет'!D5:D8)"),
        ("Общий предел", "=SUM('Помесячный бюджет'!F5:F8)"),
        ("Коммерческие возможности", "=SUM('KPI и рубежи'!D5:D8)"),
        ("Новая квалифицированная воронка, тонн", "=SUM('KPI и рубежи'!E5:E8)"),
    ]
    for r, (label, formula) in enumerate(values, 4):
        ws.cell(r, 1, label); ws.cell(r, 2, formula)
        if r <= 7: ws.cell(r, 2).number_format = MONEY
    ws["A11"] = "Принцип финансирования"
    ws["A11"].font = Font(bold=True, color=NAVY, size=12)
    ws.merge_cells("A12:H12")
    ws["A12"] = "450 → 650 → 850 → 1 050 тыс. ₸. Следующая ступень открывается только после проверки Битрикс24, качества возможностей и производственных ограничений."
    ws["A12"].fill = PatternFill("solid", fgColor=GREEN)
    ws["A12"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[12].height = 42
    ws["A14"] = "Ключевые ограничения"
    ws["A14"].font = Font(bold=True, color=NAVY, size=12)
    constraints = [
        "Лид считается квалифицированным только при наличии компании, объекта, региона, тоннажа, стадии, ЛПР и следующего шага.",
        "После 100 000 ₸ расхода без квалифицированной возможности кампания останавливается.",
        "Отраслевые СМИ запускаются только с отдельной посадочной страницей и измеряемой воронкой.",
        "Маркетинг не обещает срок и объём без подтверждённого производственного слота.",
    ]
    for i, item in enumerate(constraints, 15):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        ws.cell(i, 1, "• " + item).alignment = Alignment(wrap_text=True)
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Постепенное увеличение общего бюджета"
    chart.y_axis.title = "тенге"
    chart.x_axis.title = "месяц"
    chart.add_data(Reference(wb["Помесячный бюджет"], min_col=6, min_row=4, max_row=8), titles_from_data=True)
    chart.set_categories(Reference(wb["Помесячный бюджет"], min_col=1, min_row=5, max_row=8))
    chart.height = 7
    chart.width = 13
    ws.add_chart(chart, "D3")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 20
    for col in "CDEFGH": ws.column_dimensions[col].width = 14
    ws.freeze_panes = "A3"


def build_monthly(wb):
    ws = wb.create_sheet("Помесячный бюджет")
    title(ws, "Помесячный предел бюджета", 8)
    ws.append([]); ws.append([])
    ws.append(["Месяц", "Операционный маркетинг", "Инфраструктура", "Разовое оборудование", "Рост операционного бюджета", "Общий предел", "Факт", "Отклонение"])
    header(ws[4])
    for idx, month in enumerate(MONTHS, 5):
        ws.cell(idx, 1, month)
        ws.cell(idx, 2, MONTHLY_OPERATING[month])
        ws.cell(idx, 3, MANDATORY_INFRA[month])
        ws.cell(idx, 4, ONE_TIME_CAPEX[month])
        ws.cell(idx, 5, "—" if idx == 5 else f"=B{idx}/B{idx-1}-1")
        ws.cell(idx, 6, f"=SUM(B{idx}:D{idx})")
        ws.cell(idx, 7, None)
        ws.cell(idx, 8, f'=IF(G{idx}="","",G{idx}-F{idx})')
    ws.append(["Итого", "=SUM(B5:B8)", "=SUM(C5:C8)", "=SUM(D5:D8)", "", "=SUM(F5:F8)", "=SUM(G5:G8)", '=IF(G9=0,"",G9-F9)'])
    for row in ws.iter_rows(min_row=5, max_row=9, min_col=2, max_col=8):
        for c in row:
            if c.column != 5: c.number_format = MONEY
    for c in ws[9]: c.font = Font(bold=True)
    ws["E6"].number_format = PERCENT; ws["E7"].number_format = PERCENT; ws["E8"].number_format = PERCENT
    ws["G4"].comment = Comment("Вносится ответственным после закрытия месяца.", "Hermes")
    ws.conditional_formatting.add("H5:H9", CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))
    ws.conditional_formatting.add("H5:H9", CellIsRule(operator="lessThanOrEqual", formula=["0"], fill=PatternFill("solid", fgColor=GREEN)))
    body_style(ws, 5, 9, 1, 8)
    add_table(ws, "A4:H9", "MonthlyBudget")
    widths = [16, 23, 19, 21, 20, 18, 18, 18]
    for i, w in enumerate(widths, 1): ws.column_dimensions[chr(64+i)].width = w
    ws.freeze_panes = "A5"


def build_detail(wb):
    ws = wb.create_sheet("Детальный план")
    title(ws, "Детальный план расходов и план/факт", 10)
    ws.append([]); ws.append([])
    ws.append(["Месяц", "Статья", "План", "Факт", "Отклонение", "Ответственный", "Ожидаемый результат", "Условие запуска/масштабирования", "Статус", "Комментарий"])
    header(ws[4])
    owners = {
        CATEGORIES[0]: "Таргетолог + РОП", CATEGORIES[1]: "Таргетолог", CATEGORIES[2]: "SEO + веб-дизайнер",
        CATEGORIES[3]: "Маркетинг + технический эксперт", CATEGORIES[4]: "Мобилограф", CATEGORIES[5]: "Коммерческий директор + РОП",
        CATEGORIES[6]: "Веб-дизайнер + аналитик", CATEGORIES[7]: "Маркетинг", CATEGORIES[8]: "Коммерческий директор", CATEGORIES[9]: "Коммерческий директор",
    }
    outcomes = {
        CATEGORIES[0]: "Квалифицированные проектные обращения", CATEGORIES[1]: "Возврат посетителей отраслевых страниц", CATEGORIES[2]: "Рост видимости и доверия доменов",
        CATEGORIES[3]: "Кейсы, статьи и лид-магниты", CATEGORIES[4]: "Доказательный видеоконтент", CATEGORIES[5]: "Целевые встречи, партнёры и объекты",
        CATEGORIES[6]: "Корректная атрибуция и конверсия страниц", CATEGORIES[7]: "Ускорение производства материалов", CATEGORIES[8]: "Проверяемый доступ к отраслевым ЛПР", CATEGORIES[9]: "Проверка новых гипотез",
    }
    gates = {
        CATEGORIES[0]: "≥1 квалифицированная возможность на 100 000 ₸", CATEGORIES[1]: "Есть аудитория ≥300 посетителей", CATEGORIES[2]: "Аудит домена и перечень проверенных площадок",
        CATEGORIES[3]: "Фактология проверена техническим экспертом", CATEGORIES[4]: "Сценарий и объект согласованы", CATEGORIES[5]: "Список целевых компаний и следующий шаг",
        CATEGORIES[6]: "События и UTM перед запуском рекламы", CATEGORIES[7]: "Человеческая проверка каждого материала", CATEGORIES[8]: "Посадочная страница, UTM и отслеживаемый номер", CATEGORIES[9]: "Письменная гипотеза и критерий остановки",
    }
    row = 5
    for month in MONTHS:
        for category, amount in ALLOCATION[month].items():
            ws.append([month, category, amount, None, f'=IF(D{row}="","",D{row}-C{row})', owners[category], outcomes[category], gates[category], "Запланировано", ""])
            ws.cell(row, 3).number_format = MONEY; ws.cell(row, 4).number_format = MONEY; ws.cell(row, 5).number_format = MONEY
            row += 1
    dv = DataValidation(type="list", formula1='"Запланировано,Согласовано,В работе,Завершено,Остановлено"')
    ws.add_data_validation(dv); dv.add(f"I5:I{row-1}")
    body_style(ws, 5, row-1, 1, 10)
    add_table(ws, f"A4:J{row-1}", "DetailedBudget")
    for i, w in enumerate([14,34,15,15,15,22,35,42,17,30],1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes = "A5"; ws.auto_filter.ref = f"A4:J{row-1}"


def build_kpi(wb):
    ws = wb.create_sheet("KPI и рубежи")
    title(ws, "Показатели и условия перехода к следующей ступени", 11)
    ws.append([]); ws.append([])
    ws.append(["Месяц", "Все обращения", "Маркетингово квалифицированные", "Коммерческие возможности", "Новая воронка, тонн", "Стоимость возможности", "Стоимость тонны воронки", "Заполнение Битрикс24", "Срок реакции", "Условие перехода", "Решение"])
    header(ws[4])
    gates = {
        "Сентябрь": "Настроена аналитика; ≥2 возможностей; 95% карточек заполнены",
        "Октябрь": "≥3 возможностей; есть первая атрибуция каналов; производство подтверждает приоритетные сегменты",
        "Ноябрь": "≥4 возможностей; стоимость возможности ≤225 тыс. ₸; отраслевой пилот измерим",
        "Декабрь": "≥5 возможностей; сформирована воронка 300 т; утверждён план I квартала 2027",
    }
    for r, month in enumerate(MONTHS, 5):
        k = KPI_TARGETS[month]
        ws.append([month, k["inquiries"], k["mql"], k["sql"], k["pipeline_tons"], f"='Помесячный бюджет'!B{r}/D{r}", f"='Помесячный бюджет'!B{r}/E{r}", 0.95, "≤1 рабочего дня", gates[month], "Проверить по факту"])
        ws.cell(r,6).number_format=MONEY; ws.cell(r,7).number_format=MONEY; ws.cell(r,8).number_format=PERCENT
    body_style(ws,5,8,1,11)
    add_table(ws,"A4:K8","KPIGates")
    for i,w in enumerate([14,14,22,20,18,20,19,18,18,52,22],1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A5"


def build_audit(wb):
    ws=wb.create_sheet("Аудит исходника")
    title(ws,"Аудит файла «Маркетинговый бюджет.xlsx»",6)
    ws.append([]);ws.append([])
    ws.append(["№","Наблюдение","Риск","Рекомендация","Приоритет","Статус"]);header(ws[4])
    rows=[
        (1,"Три медиаплана: 235 216 / 482 583 / 708 820 ₸","Это только платный трафик, без полного маркетингового контура","Использовать как исходную гипотезу, заменить фактом Битрикс24","Высокий"),
        (2,"Прогноз 59–170 лидов","Не определено качество лида и отсутствует фактическая конверсия","Считать коммерческие возможности и квалифицированные тонны","Критический"),
        (3,"Аудитория M-W 25–55, Казахстан кроме запада","Слишком широкая сегментация для проектного B2B","Сегментировать по отраслям, объектам, ролям и регионам","Высокий"),
        (4,"Google/Яндекс рассчитаны в USD по курсу 488","Не учтены курс, НДС, комиссия и сопровождение","Вести утверждённый предел в тенге; комиссии — отдельными строками","Средний"),
        (5,"74 000 ₸ на внешние ссылки","Неясно: на один сайт или на оба; нет перечня площадок","Начать с пилота 37 000 ₸ и масштабировать после аудита","Высокий"),
        (6,"ElDala: баннеры, статья, новость","Охват площадки не равен квалифицированным объектам","Только пилот с UTM, посадочной страницей и стоп-критерием","Высокий"),
        (7,"Телефон от 180 000 ₸","Нет технического задания и критерия выбора","Утвердить спецификацию и лимит до покупки","Средний"),
        (8,"Зарплата мобилографа не указана","Нельзя определить полный бюджет функции","Вынести фонд оплаты труда отдельным решением","Высокий"),
        (9,"Нет план/факта и правил остановки","Расходы могут расти без результата","Ежемесячный комитет и контрольные рубежи","Критический"),
        (10,"Нет связи с производственными слотами","Маркетинг может создать неисполнимый спрос","Синхронизировать сегменты и обещания с производством","Критический"),
    ]
    for item in rows: ws.append([*item,"Открыто"])
    body_style(ws,5,14,1,6);add_table(ws,"A4:F14","SourceAudit")
    for i,w in enumerate([7,42,42,52,15,15],1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A5"


def build_content(wb):
    ws=wb.create_sheet("Контент-план")
    title(ws,"Контент и доказательства для поддержки бюджета",8)
    ws.append([]);ws.append([])
    ws.append(["Месяц","Материал","Формат","Целевой сегмент","Ответственный","Призыв к действию","Показатель","Статус"]);header(ws[4])
    items=[
        ("Сентябрь","ЛСТК против традиционного строительства: честный разбор","Статья + карусель","Промышленность/склады","Маркетинг + инженер","Получить предварительный разбор объекта","Квалифицированные обращения","В планах"),
        ("Сентябрь","Гид: 7 вопросов заводу перед договором","Лид-магнит","Все B2B","Маркетинг + РОП","Скачать после формы с объектом и тоннажем","Заполненные формы","В планах"),
        ("Сентябрь","Экскурсия на производство","Видео + очное событие","Генподрядчики/проектировщики","Мобилограф + РОП","Записаться на техническую встречу","Участники и встречи","В планах"),
        ("Октябрь","Сделано в Казахстане: 5 причин","Статья + видео","Агро/промышленность","Маркетинг","Запросить расчёт","Возможности","В планах"),
        ("Октябрь","С-профиль, Сигма, Z или П-профиль","Технический разбор","Проектировщики","Инженер + маркетинг","Запросить узлы","Проектные контакты","В планах"),
        ("Октябрь","Стройка зимой","Видео-кейс","Склады/агро","Мобилограф + производство","Обсудить зимний график","Встречи","В планах"),
        ("Ноябрь","Быстровозводимые здания для сельского хозяйства","Отраслевая страница","Агропром","Веб-дизайнер + КД","Получить отраслевое предложение","Тонны воронки","В планах"),
        ("Ноябрь","Точка росы и конденсат","Вебинар","Склады/агро","Инженер + РОП","Записаться на консультацию","SQL","В планах"),
        ("Ноябрь","Реальный кейс экономии на фундаменте","Кейс с проверенными цифрами","Собственники объектов","КД + инженер","Получить сравнение вариантов","Возможности","В планах"),
        ("Декабрь","Каркас + сэндвич-панели","Комплект решения","Промышленность/логистика","Коммерческий блок","Запросить тёплый контур","Средний тоннаж","В планах"),
        ("Декабрь","Сейсмостойкость ЛСТК","Технический материал","Юг Казахстана","Технический блок","Техническая встреча","Проектные контакты","В планах"),
        ("Декабрь","Итоги года: производство, люди, объекты","Документальный ролик","Все сегменты","Мобилограф + КД","Запланировать объект 2027","Воронка I кв. 2027","В планах"),
    ]
    for x in items: ws.append(x)
    body_style(ws,5,16,1,8);add_table(ws,"A4:H16","ContentPlan")
    for i,w in enumerate([14,48,22,25,27,35,22,15],1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A5"


def build_assumptions(wb):
    ws=wb.create_sheet("Допущения")
    title(ws,"Допущения, исключения и вопросы для утверждения",4)
    ws.append([]);ws.append([])
    ws.append(["Параметр","Принято в модели","Что требуется уточнить","Влияние"]);header(ws[4])
    rows=[
        ("Период","Сентябрь–декабрь 2026","—","План на четыре месяца"),
        ("Операционный бюджет","3 000 000 ₸","Утверждение Генерального директора","Предел расходов"),
        ("Инфраструктура","Хостинг 7 100 ₸/мес.; домены по датам исходника","Подтвердить факт уже произведённых оплат","Может снизить остаток"),
        ("Телефон мобилографа","До 180 000 ₸","Техническое задание и модель","Разовый капвклад"),
        ("Фонд оплаты труда","Не включён","Зарплаты и налоги действующего/нового штата","Добавляется сверх модели"),
        ("НДС и комиссии","Не включены","Условия подрядчиков и рекламных кабинетов","Добавляются отдельной строкой"),
        ("SEO-ссылки","37 тыс. ₸ пилот; затем 74–100 тыс. ₸","74 тыс. ₸ относится к одному или двум сайтам","Может изменить SEO-бюджет"),
        ("Курс USD/KZT","В новой модели лимиты фиксируются в тенге","Курс оплаты зарубежных сервисов","Поглощается резервом до его предела"),
        ("Лиды исходника","Не используются как гарантированный прогноз","Фактические данные Битрикс24 за 2 месяца","Заменят целевые KPI"),
        ("Производственная мощность","500 т/мес. — целевое состояние","Фактическая мощность, слоты, персонал","Ограничивает масштабирование рекламы"),
    ]
    for x in rows: ws.append(x)
    body_style(ws,5,14,1,4);add_table(ws,"A4:D14","Assumptions")
    for i,w in enumerate([28,48,48,36],1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A5"


def build_docx():
    doc=Document()
    sec=doc.sections[0];sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.5)
    doc.styles["Normal"].font.name="Arial";doc.styles["Normal"].font.size=Pt(10.5)
    for name,size,color in (("Title",22,"123047"),("Heading 1",16,"123047"),("Heading 2",13,"0F766E")):
        doc.styles[name].font.name="Arial";doc.styles[name].font.size=Pt(size);doc.styles[name].font.color.rgb=RGBColor.from_string(color)
    p=doc.add_paragraph("Маркетинговый бюджет до конца 2026 года",style="Title");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph("ТОО «Стар Билдинг» и Grand Construction");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Решение для утверждения",level=1)
    doc.add_paragraph("Утвердить постепенный операционный маркетинговый бюджет 3 000 000 тенге на сентябрь–декабрь 2026 года, обязательную цифровую инфраструктуру 49 260 тенге и разовое оборудование до 180 000 тенге. Общий предел — 3 229 260 тенге.")
    doc.add_heading("Помесячный предел",level=1)
    table=doc.add_table(rows=1,cols=5);table.style="Table Grid"
    for c,v in zip(table.rows[0].cells,["Месяц","Маркетинг","Инфраструктура","Оборудование","Всего"]): c.text=v
    for month in MONTHS:
        cells=table.add_row().cells
        vals=[month,MONTHLY_OPERATING[month],MANDATORY_INFRA[month],ONE_TIME_CAPEX[month],monthly_total(month)]
        for c,v in zip(cells,vals): c.text=f"{v:,.0f} ₸".replace(","," ") if isinstance(v,int) else v
    doc.add_heading("Логика постепенного роста",level=1)
    for text in [
        "Сентябрь: аналитика, посадочные страницы, пилот поискового спроса, квалификация в Битрикс24.",
        "Октябрь: увеличение только работающих кампаний, полный SEO-пилот, отраслевой контент и партнёрства.",
        "Ноябрь: расширение адресного маркетинга и первый измеримый пилот отраслевого СМИ.",
        "Декабрь: масштабирование подтверждённых каналов и формирование воронки на I квартал 2027 года.",
    ]: doc.add_paragraph(text,style="List Bullet")
    doc.add_heading("Контрольные показатели",level=1)
    table=doc.add_table(rows=1,cols=5);table.style="Table Grid"
    for c,v in zip(table.rows[0].cells,["Месяц","Обращения","MQL","Возможности","Воронка, т"]): c.text=v
    for month in MONTHS:
        k=KPI_TARGETS[month];cells=table.add_row().cells
        for c,v in zip(cells,[month,k['inquiries'],k['mql'],k['sql'],k['pipeline_tons']]): c.text=str(v)
    doc.add_paragraph("Показатели являются контрольными ориентирами, а не гарантированным прогнозом. Через два месяца они заменяются фактическими конверсиями Битрикс24.")
    doc.add_heading("Условия масштабирования",level=1)
    for text in [
        "Заполнение обязательных полей Битрикс24 — не менее 95%.",
        "Стоимость коммерческой возможности — не более 225 000 тенге с целевым снижением до 210 000 тенге.",
        "Кампания останавливается после 100 000 тенге расходов без квалифицированной возможности.",
        "Отраслевое СМИ запускается только с UTM, посадочной страницей и отслеживаемым обращением.",
        "Маркетинг не обещает клиенту срок и объём без подтверждения производственного слота.",
    ]: doc.add_paragraph(text,style="List Bullet")
    doc.add_heading("Что не включено",level=1)
    doc.add_paragraph("Фонд оплаты труда, НДС, агентские комиссии и неуточнённые подрядные ставки. Эти суммы добавляются отдельными строками после получения исходных данных.")
    doc.add_heading("Вопросы для решения",level=1)
    for text in [
        "74 000 тенге на внешние ссылки — на один сайт или на оба?",
        "Какие зарплаты и налоги маркетинговой команды относятся к бюджету функции?",
        "Включены ли НДС и комиссии в предложения подрядчиков?",
        "Телефон мобилографа: какая утверждённая спецификация и обязателен ли расход в сентябре?",
    ]: doc.add_paragraph(text,style="List Number")
    doc.save(DOCX)


def main():
    wb=Workbook()
    wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.calculation.calcMode="auto"
    build_monthly(wb)
    build_kpi(wb)
    build_dashboard(wb)
    # reorder dashboard created as active last to first
    dashboard = wb["Панель"]
    wb._sheets.remove(dashboard)
    wb._sheets.insert(0, dashboard)
    build_detail(wb);build_audit(wb);build_content(wb);build_assumptions(wb)
    # final required order
    order=["Панель","Помесячный бюджет","Детальный план","KPI и рубежи","Аудит исходника","Контент-план","Допущения"]
    wb._sheets=[wb[n] for n in order]
    wb.active=0
    wb.save(XLSX)
    build_docx()
    print({"xlsx":str(XLSX),"docx":str(DOCX),**totals()})


if __name__ == "__main__":
    main()
