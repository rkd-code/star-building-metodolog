import unittest
from pathlib import Path

from openpyxl import load_workbook

from budget_model import MONTHLY_OPERATING, MANDATORY_INFRA, ONE_TIME_CAPEX, totals

ROOT = Path(__file__).resolve().parent


class BudgetModelTests(unittest.TestCase):
    def test_progressive_operating_budget(self):
        values = list(MONTHLY_OPERATING.values())
        self.assertEqual([450_000, 650_000, 850_000, 1_050_000], values)
        self.assertEqual(values, sorted(values))
        self.assertEqual(3_000_000, sum(values))

    def test_total_budget(self):
        result = totals()
        self.assertEqual(49_260, result["mandatory_infrastructure"])
        self.assertEqual(180_000, result["one_time_capex"])
        self.assertEqual(3_229_260, result["grand_total"])

    def test_every_month_balances(self):
        from budget_model import ALLOCATION
        for month, target in MONTHLY_OPERATING.items():
            self.assertEqual(target, sum(ALLOCATION[month].values()), month)

    def test_workbook_outputs(self):
        path = ROOT / "marketing_budget_sep_dec_2026.xlsx"
        self.assertTrue(path.exists())
        wb = load_workbook(path, data_only=False)
        self.assertEqual(
            ["Панель", "Помесячный бюджет", "Детальный план", "KPI и рубежи", "Аудит исходника", "Контент-план", "Допущения"],
            wb.sheetnames,
        )
        self.assertGreaterEqual(len(wb["Панель"]._charts), 1)
        self.assertEqual("=SUM('Помесячный бюджет'!F5:F8)", wb["Панель"]["B7"].value)


if __name__ == "__main__":
    unittest.main()
